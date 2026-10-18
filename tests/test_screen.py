"""Bulk screening (analytics.rh_screen): estimates, tiers, incremental refresh, filters, shortlists with diff, exports, portfolio maths. Uses a tiny bbox and removes its rows afterwards (make screen refills them)."""
import io, math, os
import pytest
from fastapi.testclient import TestClient
from backend.app import portfolio, screen
from backend.app.db import conn
BB = (1750000, 5916000, 1750400, 5916300)
FULL = {"fee_per_bed_per_week": 1600, "food_per_occupied_bed_per_week": 200, "staff_wage_per_week": 800, "staff_positions_per_8_beds": 4, "clinical_lead_cost_per_week": 2000, "manager_cost_per_week": 2500, "cook_cleaner_laundry_per_week": 1500,
        "utilities_per_week": 500, "insurance_per_year": 26000, "supplies_per_occupied_bed_per_week": 40, "repairs_maintenance_per_year": 26000, "compliance_and_certification_per_year": 13000, "other_opex_per_week": 300, "roster_coverage_multiplier": 1.5,
        "stabilised_occupancy": 0.9, "fill_ramp_by_month": [0.3, 0.6, 0.8], "pre_opening_costs": 20000, "working_capital": 15000, "lease_term_years": 10, "cost_of_capital_pct": 8, "target_profit_per_year": 20000, "minimum_acceptable_profit_per_year": 10000, "retrofit_cost_per_bed_screening": 10000}

@pytest.fixture(scope="module")
def client():
    from backend.app.main import app
    def wipe():
        with conn() as c:
            c.execute("DELETE FROM analytics.rh_screen WHERE lon IS NOT NULL AND parcel_id IN (SELECT parcel_id FROM core.sites WHERE geom && ST_MakeEnvelope(%s,%s,%s,%s,2193))", BB)
            c.execute("DELETE FROM core.user_shortlist"); c.execute("DELETE FROM core.user_listing WHERE source_id LIKE 'TSC-%'"); c.commit()
    wipe(); cl = TestClient(app); cl.put("/api/settings/assumptions", json={"values": {k: v for k, v in FULL.items()}}); yield cl; wipe()
    with conn() as c: c.execute("DELETE FROM core.user_setting"); c.execute("DELETE FROM core.user_setting_history"); c.commit()

def test_candidate_zones_and_estimates_hand_checked(client):
    z = screen.candidate_zones(); assert "Residential - Mixed Housing Urban Zone" in z and len(z) >= 10
    r = screen.run(BB, progress=None); assert r["seen"] > 10 and r["written"] == r["seen"] and r["uses_placeholders"] in (True, False)
    with conn() as c: rows = c.execute("SELECT * FROM analytics.rh_screen WHERE parcel_id IN (SELECT parcel_id FROM core.sites WHERE geom && ST_MakeEnvelope(%s,%s,%s,%s,2193))", BB).fetchall()
    for x in rows:
        assert x["beds_tier"] == "ESTIMATED" and x["beds_base"] == int(float(x["footprint_m2"]) * 1 // 60) and "placeholder" in x["beds_source"]
        assert x["inputs_hash"]
        if x["beds_base"] >= 1: assert float(x["retrofit_est"]) == x["beds_base"] * 10000 and x["profit_before_rent_annual"] is not None and x["walk_away_rent"] is not None
        else: assert x["profit_before_rent_annual"] is None          # no beds: no profit figure
    big = max(rows, key=lambda x: x["beds_base"]); assert big["too_small"] is (big["beds_base"] < big["min_beds_minimum"])

def test_incremental_refresh_and_blank_inputs_give_no_numbers(client):
    r2 = screen.run(BB, progress=None); assert r2["written"] == 0 and r2["unchanged"] == r2["seen"]                                  # nothing changed: nothing rewritten
    client.put("/api/settings/assumptions", json={"values": {"utilities_per_week": None}})
    r3 = screen.run(BB, progress=None); assert r3["written"] == r3["seen"]                                                      # an input changed: every row is recomputed
    with conn() as c: x = c.execute("SELECT profit_before_rent_annual, walk_away_rent, needs_input FROM analytics.rh_screen WHERE beds_base >= 1 AND parcel_id IN (SELECT parcel_id FROM core.sites WHERE geom && ST_MakeEnvelope(%s,%s,%s,%s,2193)) LIMIT 1", BB).fetchone()
    assert x["profit_before_rent_annual"] is None and x["walk_away_rent"] is None and "utilities_per_week" in x["needs_input"]
    client.put("/api/settings/assumptions", json={"values": {"utilities_per_week": 500}}); screen.run(BB, progress=None)

def test_listing_and_property_record_upgrade_the_tier(client):
    with conn() as c: p = c.execute("SELECT s.parcel_id, s.address FROM analytics.rh_screen s WHERE s.address IS NOT NULL AND s.parcel_id IN (SELECT parcel_id FROM core.sites WHERE geom && ST_MakeEnvelope(%s,%s,%s,%s,2193)) LIMIT 1", BB).fetchone()
    csv = "source_id,address,suburb,rent_per_week,bedrooms,bathrooms,listing_type,url\nTSC-1," + p["address"].split(",")[0] + ",,900,9,3,rent,\n"
    client.post("/api/listings/import", files={"file": ("l.csv", csv.encode())}, data={"source": "csv_manual", "dry_run": "false"})
    screen.run(BB, progress=None)
    with conn() as c: x = c.execute("SELECT beds_base, beds_tier, beds_source FROM analytics.rh_screen WHERE parcel_id=%s", (p["parcel_id"],)).fetchone()
    assert (x["beds_base"], x["beds_tier"]) == (9, "DOCUMENTED") and "listing" in x["beds_source"]
    f = client.post("/api/screen/query", json={"filters": {"tiers": ["DOCUMENTED"], "min_beds": 9}}).json(); assert f["total"] >= 1 and all(r["beds_tier"] == "DOCUMENTED" for r in f["rows"])

def test_query_filters_sort_polygon_radius_points(client):
    q = client.post("/api/screen/query", json={"filters": {"min_beds": 1}, "sort": "beds_base", "limit": 5}).json(); assert len(q["rows"]) <= 5 and "ESTIMATED" in q["banner"] and q["rows"] == sorted(q["rows"], key=lambda r: -r["beds_base"])
    assert client.post("/api/screen/query", json={"sort": "drop table"}).status_code == 422 and client.post("/api/screen/query", json={"filters": {"tiers": ["MAYBE"]}}).status_code == 422
    one = q["rows"][0]; lon, lat = one["lon"], one["lat"]
    assert any(r["parcel_id"] == one["parcel_id"] for r in client.post("/api/screen/query", json={"filters": {"radius": [lat, lon, 0.05]}}).json()["rows"])
    ring = [[lon - .001, lat - .001], [lon + .001, lat - .001], [lon + .001, lat + .001], [lon - .001, lat + .001], [lon - .001, lat - .001]]
    assert any(r["parcel_id"] == one["parcel_id"] for r in client.post("/api/screen/query", json={"filters": {"polygon": ring}}).json()["rows"]) and client.post("/api/screen/query", json={"filters": {"polygon": [[0, 0]]}}).status_code == 422
    pts = client.post("/api/screen/points", json={"bbox": [lon - .01, lat - .01, lon + .01, lat + .01]}).json(); assert pts["points"] and len(pts["points"][0]) == 7

def test_suburbs_notes_blank_not_guessed(client):
    s = client.get("/api/screen/suburbs").json(); assert s["suburbs"] and all(x["certified_beds_existing"] is None and x["beds_per_1000_aged_75_plus"] is None for x in s["suburbs"]) and any("HealthCERT" in n for n in s["notes"])

def test_parcel_explain_and_non_candidates(client):
    assert client.get("/api/screen/parcel/1").status_code == 404
    with conn() as c: nb = c.execute("SELECT parcel_id FROM core.sites WHERE building_count = 0 AND zone_name = 'Residential - Mixed Housing Urban Zone' LIMIT 1").fetchone()
    e = client.get(f"/api/screen/parcel/{nb['parcel_id']}").json(); assert e["candidate"] is False and "no building" in e["reason"]
    st = client.get("/api/screen/status").json(); assert st["rows"] > 0 and st["assumptions"]["gfa_per_bed_m2"]["flag"] == "placeholder"

def test_shortlists_diff_audit_export(client):
    sl = client.post("/api/shortlists", json={"name": "big ones", "filters": {"min_beds": 2, "exclude_too_small": True}, "notes": "test"}).json()
    r1 = client.post(f"/api/shortlists/{sl['id']}/run").json(); assert r1["diff"] is None and r1["rows"]
    with conn() as c: c.execute("DELETE FROM analytics.rh_screen WHERE parcel_id=%s", (r1["rows"][0]["parcel_id"],)); c.commit()          # one row disappears
    r2 = client.post(f"/api/shortlists/{sl['id']}/run").json(); assert r1["rows"][0]["parcel_id"] in r2["diff"]["dropped"] and r1["rows"][0]["parcel_id"] in r2["diff"]["dropped_reasons"] or r2["diff"]["dropped"]
    assert [a["action"] for a in client.get(f"/api/shortlists/{sl['id']}/audit").json()] == ["created", "run", "run"]
    csv = client.get(f"/api/shortlists/{sl['id']}/export").text; assert csv.splitlines()[0].startswith("parcel_id,address") and "Screening on ESTIMATED" in csv
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/shortlists/{sl['id']}/export", params={"fmt": "xlsx"}).content)); assert wb.active.cell(1, 1).value == "parcel_id"
    assert client.post("/api/shortlists", json={"name": "", "filters": {}}).status_code == 422

def test_portfolio_hand_calculated_and_needs_input(client):
    from backend.app import economics as E
    homes = [{"parcel_id": i, "address": f"h{i}", "beds": 16, "tier": "ESTIMATED", "retrofit": 160000} for i in (1, 2, 3)]
    assert portfolio.evaluate(homes, {"stabilised_occupancy": None}, {"cash_limit_dollars": 1e6})["status"] == "needs_input"
    assert portfolio.evaluate(homes, {}, {})["status"] == "needs_input"                                        # no cash limit: no roll-out figures
    p = {**FULL}; one = portfolio.home_economics(16, 160000, p)
    r = portfolio.evaluate(homes, p, {"cash_limit_dollars": 2_000_000, "homes_per_manager": 3, "certification_cost_per_site": 5000})
    assert r["status"] == "ok" and r["homes"] == 3 and r["beds"] == 48 and r["retrofit_total"] == 480000 and r["certification_cost_total"] == 15000
    assert r["shared_staff_saving_per_year"] == pytest.approx((2000 + 2500) * 2 * E.W) and r["staff_to_hire"]["managers"] == 1
    assert r["net_operating_profit_per_year"] == pytest.approx((3 * one["net_week"] + 9000) * E.W) and r["revenue_per_year"] == pytest.approx(3 * one["revenue_week"] * E.W)
    assert r["staff_to_hire"]["care_staff_positions"] == pytest.approx(3 * 16 / 8 * 4 * 1.5) and "HealthCERT" in r["needs_professional"][0] and r["sensitivity_fill_minus_10pts"]["change"] < 0
    tight = portfolio.evaluate(homes, p, {"cash_limit_dollars": 250_000, "staff_positions_limit": 20}); assert sum(1 for t in tight["timeline_stagger"] if t["status"] != "not started within horizon") <= 1
